#!/usr/bin/env python3
"""Feasibility study financial model.

Usage: python3 fin_model.py assumptions.json [output.xlsx] > model-output.md

Reads a filled assumptions file (see templates/assumptions-template.json),
builds a 60-month model, and prints a Markdown report: startup costs, funding,
5-year income statement, cash flow, break-even, NPV, IRR, payback and a
sensitivity table. Optionally writes the same tables to an Excel workbook.
All arithmetic lives here so agents never compute financials by hand.
"""
import copy
import json
import sys


def npv(rate, flows):
    return sum(f / (1 + rate) ** i for i, f in enumerate(flows))


def irr(flows):
    lo, hi = -0.99, 10.0
    f_lo, f_hi = npv(lo, flows), npv(hi, flows)
    if f_lo * f_hi > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        f_mid = npv(mid, flows)
        if f_lo * f_mid <= 0:
            hi = mid
        else:
            lo, f_lo = mid, f_mid
    return (lo + hi) / 2


def loan_payment(principal, annual_rate, years):
    if principal <= 0:
        return 0.0
    r, n = annual_rate / 12, years * 12
    return principal / n if r == 0 else principal * r / (1 - (1 + r) ** -n)


def run(a):
    years = 5
    months = years * 12
    capex = sum(i["amount"] for i in a["capex"])
    pre_open = sum(i["amount"] for i in a.get("pre_opening_costs", []))
    annual_dep = sum(i["amount"] / i["life_years"] for i in a["capex"])

    infl = a.get("opex_inflation", 0.0)
    sal_g = a.get("salary_growth", 0.0)
    price_g = a.get("price_growth", 0.0)
    comm = a.get("channel_commission_pct", 0.0)
    emp = a.get("employer_contribution_rate", 0.0)
    mk = a["marketing"]

    rows = []
    for m in range(months):
        y = m // 12
        rev = cogs = units_total = 0.0
        for s in a["revenue_streams"]:
            units = s["units_month_y1"][m] if y == 0 else s["units_year"][y] / 12
            price = s["price"] * (1 + price_g) ** y
            ucost = s["unit_cost"] * (1 + infl) ** y
            rev += units * price
            cogs += units * ucost
            units_total += units
        commission = rev * comm
        fixed = sum(i["amount"] for i in a["fixed_opex_monthly"]) * (1 + infl) ** y
        payroll = sum(p["count_by_year"][y] * p["monthly_salary"] for p in a["payroll"]) * (1 + sal_g) ** y * (1 + emp)
        marketing = mk["monthly_y1"] if y == 0 else rev * mk["pct_revenue_after_y1"]
        ebitda = rev - cogs - commission - fixed - payroll - marketing
        rows.append(dict(month=m + 1, year=y + 1, units=units_total, revenue=rev, cogs=cogs,
                         commission=commission, fixed=fixed, payroll=payroll,
                         marketing=marketing, ebitda=ebitda))

    # Working capital reserve: N months of year-1 average cash costs
    y1_cash_costs = sum(r["fixed"] + r["payroll"] + r["marketing"] for r in rows[:12]) / 12
    wc = a.get("working_capital_months", 3) * y1_cash_costs
    startup = capex + pre_open + wc

    fin = a.get("financing", {})
    loan = fin.get("loan_amount", 0.0)
    pay = loan_payment(loan, fin.get("loan_rate", 0.0), fin.get("loan_years", 1))
    bal = loan
    tax_rate = a["corporate_tax_rate"]

    annual = []
    for y in range(years):
        yr = rows[y * 12:(y + 1) * 12]
        s = {k: sum(r[k] for r in yr) for k in ("units", "revenue", "cogs", "commission", "fixed", "payroll", "marketing", "ebitda")}
        interest = principal = 0.0
        for _ in range(12):
            if bal > 1e-9:
                i = bal * fin.get("loan_rate", 0.0) / 12
                p = min(pay - i, bal)
                interest += i
                principal += p
                bal -= p
        dep = annual_dep if y < max(i["life_years"] for i in a["capex"]) else sum(
            i["amount"] / i["life_years"] for i in a["capex"] if i["life_years"] > y)
        ebt = s["ebitda"] - dep - interest - (pre_open if y == 0 else 0)
        tax = max(0.0, ebt) * tax_rate
        net = ebt - tax
        s.update(depreciation=dep, interest=interest, pre_opening=pre_open if y == 0 else 0,
                 ebt=ebt, tax=tax, net_profit=net, loan_principal=principal,
                 project_cf=s["ebitda"] - tax,
                 equity_cf=s["ebitda"] - tax - interest - principal)
        annual.append(s)

    # Recover working capital and residual book value at end of year 5
    residual = sum(max(0.0, i["amount"] * (1 - years / i["life_years"])) for i in a["capex"])
    flows = [-startup] + [s["project_cf"] for s in annual]
    terminal = wc + residual * a.get("residual_value_pct", 0.0)
    flows[-1] += terminal
    annual[-1]["equity_cf"] += terminal
    rate = a["discount_rate"]

    # Monthly payback (project basis, tax approximated at annual rate per month)
    cum, payback_m = -startup, None
    for r in rows:
        yr = annual[r["year"] - 1]
        cum += r["ebitda"] - yr["tax"] / 12
        if payback_m is None and cum >= 0:
            payback_m = r["month"]
    be_month = next((r["month"] for r in rows if r["ebitda"] >= 0), None)

    # Break-even units per month in year 1 (weighted contribution margin)
    tot_units = sum(r["units"] for r in rows[:12]) or 1
    tot_rev = sum(r["revenue"] for r in rows[:12])
    tot_var = sum(r["cogs"] + r["commission"] for r in rows[:12])
    cm_unit = (tot_rev - tot_var) / tot_units
    fixed_m = sum(r["fixed"] + r["payroll"] + r["marketing"] for r in rows[:12]) / 12
    be_units = fixed_m / cm_unit if cm_unit > 0 else None

    return dict(capex=capex, pre_open=pre_open, wc=wc, startup=startup, loan=loan,
                equity=startup - loan, rows=rows, annual=annual, flows=flows,
                npv=npv(rate, flows), irr=irr(flows), payback_m=payback_m,
                be_month=be_month, be_units=be_units, cm_unit=cm_unit,
                avg_units_y1=tot_units / 12, rate=rate)


def sensitivity(a):
    cases = [
        ("Base case", lambda x: x),
        ("Price -10%", lambda x: _scale(x, "price", 0.9)),
        ("Volume -20%", lambda x: _vol(x, 0.8)),
        ("Volume +20%", lambda x: _vol(x, 1.2)),
        ("Unit cost +15%", lambda x: _scale(x, "unit_cost", 1.15)),
        ("Fixed costs & payroll +15%", lambda x: _fixed(x, 1.15)),
        ("Capex +20%", lambda x: _capex(x, 1.2)),
        ("Pessimistic (vol -20%, cost +15%)", lambda x: _scale(_vol(x, 0.8), "unit_cost", 1.15)),
    ]
    out = []
    for name, f in cases:
        r = run(f(copy.deepcopy(a)))
        out.append((name, r["npv"], r["irr"], r["payback_m"], r["annual"][0]["net_profit"]))
    return out


def _scale(a, key, k):
    for s in a["revenue_streams"]:
        s[key] *= k
    return a


def _vol(a, k):
    for s in a["revenue_streams"]:
        s["units_month_y1"] = [u * k for u in s["units_month_y1"]]
        s["units_year"] = [u * k for u in s["units_year"]]
    return a


def _fixed(a, k):
    for i in a["fixed_opex_monthly"]:
        i["amount"] *= k
    for p in a["payroll"]:
        p["monthly_salary"] *= k
    return a


def _capex(a, k):
    for i in a["capex"]:
        i["amount"] *= k
    return a


def fmt(x, cur=""):
    if x is None:
        return "n/a"
    return f"{cur}{x:,.0f}"


def pct(x):
    return "n/a (never positive)" if x is None else f"{x * 100:.1f}%"


def report(a, r, sens):
    c = a["currency"] + " "
    L = [f"# Financial model output: {a['project']}", "",
         f"All figures in {a['currency']}. Generated by tools/fin_model.py from assumptions.json.", "",
         "## Startup cost and funding", "", "| Item | Amount |", "|---|---:|"]
    for i in a["capex"]:
        L.append(f"| {i['name']} (capex) | {fmt(i['amount'])} |")
    for i in a.get("pre_opening_costs", []):
        L.append(f"| {i['name']} (pre-opening) | {fmt(i['amount'])} |")
    L += [f"| Working capital reserve ({a.get('working_capital_months', 3)} months) | {fmt(r['wc'])} |",
          f"| **Total startup cost** | **{fmt(r['startup'])}** |",
          f"| Funded by equity | {fmt(r['equity'])} |", f"| Funded by loan | {fmt(r['loan'])} |", "",
          "## Income statement (5 years)", "",
          "| Line | " + " | ".join(f"Y{i}" for i in range(1, 6)) + " |", "|---|" + "---:|" * 5]
    lines = [("Units sold", "units"), ("Revenue", "revenue"), ("Cost of goods", "cogs"),
             ("Channel commissions", "commission"), ("Payroll (incl. employer contributions)", "payroll"),
             ("Fixed operating costs", "fixed"), ("Marketing", "marketing"), ("**EBITDA**", "ebitda"),
             ("Pre-opening costs expensed", "pre_opening"), ("Depreciation", "depreciation"),
             ("Interest", "interest"), ("Profit before tax", "ebt"), ("Tax", "tax"),
             ("**Net profit**", "net_profit")]
    for label, k in lines:
        L.append(f"| {label} | " + " | ".join(fmt(s[k]) for s in r["annual"]) + " |")
    L.append("| Net margin | " + " | ".join(pct(s["net_profit"] / s["revenue"]) if s["revenue"] else "n/a" for s in r["annual"]) + " |")
    L += ["", "## Cash flow", "", "| | Y0 | " + " | ".join(f"Y{i}" for i in range(1, 6)) + " |", "|---|" + "---:|" * 6,
          "| Project free cash flow | " + " | ".join(fmt(f) for f in r["flows"]) + " |"]
    cum, cums = 0, []
    for f in r["flows"]:
        cum += f
        cums.append(cum)
    L.append("| Cumulative | " + " | ".join(fmt(x) for x in cums) + " |")
    L.append("| Cash to owner after loan service | " + fmt(-r["equity"]) + " | " + " | ".join(fmt(s["equity_cf"]) for s in r["annual"]) + " |")
    pb = r["payback_m"]
    L += ["", "## Key indicators", "", "| Indicator | Value |", "|---|---:|",
          f"| Discount rate | {pct(r['rate'])} |",
          f"| NPV (5 years) | {c}{fmt(r['npv'])} |",
          f"| IRR (5 years) | {pct(r['irr'])} |",
          f"| Payback | {'not within 5 years' if pb is None else f'month {pb} ({pb / 12:.1f} years)'} |",
          f"| First month with positive EBITDA | {r['be_month'] or 'not within 5 years'} |",
          f"| Break-even volume (year 1) | {fmt(r['be_units'])} units/month vs plan avg {fmt(r['avg_units_y1'])} |",
          f"| Contribution margin per unit (year 1) | {c}{r['cm_unit']:,.2f} |",
          f"| 5-year ROI on startup cost | {pct(sum(s['net_profit'] for s in r['annual']) / r['startup'])} |",
          "", "## Sensitivity", "", "| Scenario | NPV | IRR | Payback month | Year-1 net profit |", "|---|---:|---:|---:|---:|"]
    for name, n, i, p, y1 in sens:
        L.append(f"| {name} | {fmt(n)} | {pct(i)} | {p or '>60'} | {fmt(y1)} |")
    L += ["", "## Year 1 by month", "", "| Month | Units | Revenue | EBITDA |", "|---:|---:|---:|---:|"]
    for row in r["rows"][:12]:
        L.append(f"| {row['month']} | {fmt(row['units'])} | {fmt(row['revenue'])} | {fmt(row['ebitda'])} |")
    return "\n".join(L)


def write_xlsx(path, a, r, sens):
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    ws.append(["Indicator", "Value"])
    for k, v in [("Total startup cost", r["startup"]), ("NPV", r["npv"]), ("IRR", r["irr"]),
                 ("Payback month", r["payback_m"]), ("First positive EBITDA month", r["be_month"]),
                 ("Break-even units/month (Y1)", r["be_units"])]:
        ws.append([k, v])
    pl = wb.create_sheet("Income statement")
    keys = ["units", "revenue", "cogs", "commission", "payroll", "fixed", "marketing", "ebitda",
            "pre_opening", "depreciation", "interest", "ebt", "tax", "net_profit", "project_cf"]
    pl.append(["Line"] + [f"Y{i}" for i in range(1, 6)])
    for k in keys:
        pl.append([k] + [round(s[k], 2) for s in r["annual"]])
    mo = wb.create_sheet("Monthly")
    cols = ["month", "year", "units", "revenue", "cogs", "commission", "fixed", "payroll", "marketing", "ebitda"]
    mo.append(cols)
    for row in r["rows"]:
        mo.append([round(row[c], 2) if isinstance(row[c], float) else row[c] for c in cols])
    se = wb.create_sheet("Sensitivity")
    se.append(["Scenario", "NPV", "IRR", "Payback month", "Year-1 net profit"])
    for s in sens:
        se.append(list(s))
    asn = wb.create_sheet("Assumptions")
    asn.append(["Assumptions (JSON)"])
    for line in json.dumps(a, indent=2, ensure_ascii=False).splitlines():
        asn.append([line])
    for sh in wb.worksheets:
        for cell in sh[1]:
            cell.font = Font(bold=True)
        sh.column_dimensions["A"].width = 34
    wb.save(path)


def main():
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    with open(sys.argv[1], encoding="utf-8") as f:
        a = json.load(f)
    r = run(a)
    sens = sensitivity(a)
    print(report(a, r, sens))
    if len(sys.argv) > 2:
        write_xlsx(sys.argv[2], a, r, sens)


if __name__ == "__main__":
    main()
